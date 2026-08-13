"""Unit tests for the Split_DOCX song splitter."""

import contextlib
import io
import os
import sys
import tempfile
import unittest
from unittest import mock

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "src"))

import Split_DOCX  # noqa: E402


class SplitSongsTests(unittest.TestCase):
    def test_splits_on_tabbed_titles(self):
        text = "\tSong One\nline1\nline2\n\tSong Two\nline3\n"
        songs = Split_DOCX.split_songs(text)
        self.assertEqual(len(songs), 2)
        self.assertEqual(songs[0][0], "Song One")
        self.assertEqual(songs[0][1], ["line1", "line2"])
        self.assertEqual(songs[1][0], "Song Two")
        self.assertEqual(songs[1][1], ["line3"])

    def test_blank_input_yields_no_songs(self):
        self.assertEqual(Split_DOCX.split_songs(""), [])

    def test_no_tabbed_titles_yields_no_songs(self):
        text = "just lyrics\nwith no tabs\nat all\n"
        self.assertEqual(Split_DOCX.split_songs(text), [])

    def test_crlf_line_endings_are_handled(self):
        text = "\tTitle\r\nlyric one\r\nlyric two\r\n"
        songs = Split_DOCX.split_songs(text)
        self.assertEqual(songs[0][0], "Title")
        self.assertEqual(songs[0][1], ["lyric one", "lyric two"])

    def test_untabbed_markers_stay_inside_a_song(self):
        text = "\tSong\n[Chorus]\nlyric\nRefrain\nmore lyric\n"
        songs = Split_DOCX.split_songs(text)
        self.assertEqual(songs[0][1], ["[Chorus]", "lyric", "Refrain", "more lyric"])

    def test_safe_filename(self):
        self.assertEqual(Split_DOCX.safe_filename("A:B? /x*"), "A_B_ _x_")

    def test_safe_filename_falls_back(self):
        self.assertEqual(Split_DOCX.safe_filename("   "), "untitled")

    def test_unique_filename(self):
        used = {}
        self.assertEqual(Split_DOCX.unique_filename("Song", "rtf", used), "Song.rtf")
        self.assertEqual(Split_DOCX.unique_filename("Song", "rtf", used), "Song (2).rtf")
        self.assertEqual(Split_DOCX.unique_filename("Song", "rtf", used), "Song (3).rtf")
        self.assertEqual(Split_DOCX.unique_filename("Other", "txt", used), "Other.txt")


class MissingInputTests(unittest.TestCase):
    def test_main_reports_missing_input(self):
        buf = io.StringIO()
        with contextlib.redirect_stdout(buf):
            rc = Split_DOCX.main(["/no/such/file.docx", "out"])
        self.assertEqual(rc, 1)
        self.assertIn("Input file not found", buf.getvalue())


class DocxTests(unittest.TestCase):
    def _build_docx(self, path, title_size_pt=14.0):
        """Build a small two-song .docx with bold upsized titles and a bold
        (but not upsized) section marker that must NOT be treated as a title."""
        from docx import Document
        from docx.shared import Pt

        doc = Document()

        p = doc.add_paragraph()
        r = p.add_run("Processional Song: Sing A New Song")
        r.bold = True
        r.font.size = Pt(title_size_pt)
        doc.add_paragraph("Sing a new song unto the Lord;")

        ref = doc.add_paragraph()
        rr = ref.add_run("[Chorus]")
        rr.bold = True  # bold but not upsized -> must not be a title
        doc.add_paragraph("Singing alleluia")

        p = doc.add_paragraph()
        r = p.add_run("Communion Song: Holy Is His Name")
        r.bold = True
        r.font.size = Pt(title_size_pt)
        doc.add_paragraph("My soul proclaims the greatness of the Lord")

        doc.save(path)

    def test_split_into_songs_picks_up_bold_upsized_titles(self):
        with tempfile.TemporaryDirectory() as tmp:
            inp = os.path.join(tmp, "songs.docx")
            self._build_docx(inp)
            songs = Split_DOCX.split_into_songs(inp)
        self.assertEqual(len(songs), 2)
        # Role prefix stripped for the clean title...
        self.assertEqual(songs[0]["title"], "Sing A New Song")
        self.assertEqual(songs[0]["raw_title"], "Processional Song: Sing A New Song")
        # ...while the bold-but-not-upsized [Chorus] stayed as a lyric line.
        self.assertEqual(
            songs[0]["body_lines"],
            ["Sing a new song unto the Lord;", "[Chorus]", "Singing alleluia"],
        )
        self.assertEqual(songs[1]["title"], "Holy Is His Name")

    def test_title_split_across_multiple_runs_is_detected(self):
        """Titles are often split across several runs by Word; the classifier
        must detect a bold+upsized run even when it isn't the first run."""
        from docx import Document
        from docx.shared import Pt

        with tempfile.TemporaryDirectory() as tmp:
            inp = os.path.join(tmp, "songs.docx")
            doc = Document()
            p = doc.add_paragraph()
            p.add_run("Processional Song: ")          # plain run first
            r2 = p.add_run("Split Across Runs")        # bold+upsize on 2nd run
            r2.bold = True
            r2.font.size = Pt(14)
            doc.add_paragraph("A lyric line")
            doc.save(inp)

            songs = Split_DOCX.split_into_songs(inp)
        self.assertEqual(len(songs), 1)
        self.assertEqual(songs[0]["title"], "Split Across Runs")
        self.assertEqual(songs[0]["body_lines"], ["A lyric line"])

    def test_processDOCX_corrupt_file_returns_error(self):
        with tempfile.TemporaryDirectory() as tmp:
            inp = os.path.join(tmp, "bad.docx")
            out = os.path.join(tmp, "out")
            with open(inp, "w", encoding="utf-8") as fh:
                fh.write("this is not a real docx")
            buf = io.StringIO()
            with contextlib.redirect_stdout(buf):
                rc = Split_DOCX.processDOCX(inp, out)
            self.assertEqual(rc, 1)
            self.assertIn("Failed to read DOCX", buf.getvalue())

    def test_processDOCX_writes_rtf_per_song(self):
        with tempfile.TemporaryDirectory() as tmp:
            inp = os.path.join(tmp, "songs.docx")
            out = os.path.join(tmp, "out")
            self._build_docx(inp)
            buf = io.StringIO()
            with contextlib.redirect_stdout(buf):
                rc = Split_DOCX.processDOCX(inp, out)
            self.assertEqual(rc, 0)
            self.assertEqual(
                sorted(os.listdir(out)), ["Holy Is His Name.rtf", "Sing A New Song.rtf"]
            )
            with open(os.path.join(out, "Sing A New Song.rtf"), encoding="utf-8") as fh:
                content = fh.read()
            self.assertIn(r"\b\fs28 Sing A New Song\b0", content)
            self.assertIn("Sing a new song unto the Lord;", content)

    def test_main_dispatches_docx_to_processDOCX(self):
        with tempfile.TemporaryDirectory() as tmp:
            inp = os.path.join(tmp, "songs.docx")
            out = os.path.join(tmp, "out")
            self._build_docx(inp)
            buf = io.StringIO()
            with contextlib.redirect_stdout(buf):
                rc = Split_DOCX.main([inp, out])
            self.assertEqual(rc, 0)
            self.assertTrue(os.path.exists(out))
            self.assertEqual(len(os.listdir(out)), 2)

    def test_processDOCX_reports_missing_dependency(self):
        # Simulate python-docx not being installed.
        real_import = __import__

        def fake_import(name, *args, **kwargs):
            if name == "docx" or name.startswith("docx."):
                raise ImportError("No module named 'docx'")
            return real_import(name, *args, **kwargs)

        buf = io.StringIO()
        with contextlib.redirect_stdout(buf):
            with mock.patch("builtins.__import__", side_effect=fake_import):
                rc = Split_DOCX.processDOCX("whatever.docx", "out")
        self.assertEqual(rc, 1)
        self.assertIn("python-docx is required", buf.getvalue())

    def test_main_writes_one_file_per_song(self):
        with tempfile.TemporaryDirectory() as tmp:
            inp = os.path.join(tmp, "songs.txt")
            out = os.path.join(tmp, "out")
            with open(inp, "w", encoding="utf-8") as fh:
                fh.write("\tAlpha\nl1\nl2\n\tBeta\nl3\n")

            self.assertEqual(Split_DOCX.main([inp, out]), 0)

            files = sorted(os.listdir(out))
            self.assertEqual(files, ["Alpha.txt", "Beta.txt"])

            with open(os.path.join(out, "Alpha.txt"), encoding="utf-8") as fh:
                content = fh.read()
            self.assertIn("Alpha", content)
            self.assertIn("l2", content)

    def test_main_txt_duplicate_titles_do_not_overwrite(self):
        with tempfile.TemporaryDirectory() as tmp:
            inp = os.path.join(tmp, "songs.txt")
            out = os.path.join(tmp, "out")
            with open(inp, "w", encoding="utf-8") as fh:
                fh.write("\tSame Title\nfirst\n\tSame Title\nsecond\n")

            self.assertEqual(Split_DOCX.main([inp, out]), 0)

            files = sorted(os.listdir(out))
            self.assertEqual(files, ["Same Title (2).txt", "Same Title.txt"])

            with open(os.path.join(out, "Same Title.txt"), encoding="utf-8") as fh:
                self.assertIn("first", fh.read())
            with open(os.path.join(out, "Same Title (2).txt"), encoding="utf-8") as fh:
                self.assertIn("second", fh.read())


class HelperTests(unittest.TestCase):
    def test_decode_text_utf8(self):
        self.assertEqual(Split_DOCX.decode_text("héllo".encode("utf-8")), "héllo")

    def test_decode_text_cp1252_fallback(self):
        # 0x92 is a right single quote in Windows-1252 (invalid as standalone UTF-8).
        self.assertEqual(Split_DOCX.decode_text(b"God\x92s"), "God\u2019s")

    def test_rtf_escape_tab_and_specials(self):
        self.assertEqual(Split_DOCX.rtf_escape("a\tb{c}\\d"), r"a\tab b\{c\}\\d")

    def test_resolve_input_uses_explicit_arg(self):
        self.assertEqual(Split_DOCX.resolve_input_path(["foo.docx"], "/tmp"), "foo.docx")

    def test_resolve_input_prefers_docx_when_present(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = os.path.join(tmp, "root")
            src = os.path.join(root, "src")
            os.makedirs(src, exist_ok=True)
            open(os.path.join(root, "songs.txt"), "w").close()
            open(os.path.join(root, "songs.docx"), "w").close()
            self.assertEqual(
                Split_DOCX.resolve_input_path([], src), os.path.join(root, "songs.docx")
            )

    def test_resolve_input_falls_back_to_txt(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = os.path.join(tmp, "root")
            src = os.path.join(root, "src")
            os.makedirs(src, exist_ok=True)
            open(os.path.join(root, "songs.txt"), "w").close()
            self.assertEqual(
                Split_DOCX.resolve_input_path([], src), os.path.join(root, "songs.txt")
            )

    def test_resolve_input_defaults_when_none_exist(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = os.path.join(tmp, "root")
            src = os.path.join(root, "src")
            os.makedirs(src, exist_ok=True)
            self.assertEqual(
                Split_DOCX.resolve_input_path([], src), os.path.join(root, "songs.txt")
            )


if __name__ == "__main__":
    unittest.main()
